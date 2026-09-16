#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1,<2"]
# ///
"""Render a Statement of Work DOCX from a derived format spec.

The module exposes ``SOWDoc`` for agent-authored documents and a JSON-driven CLI
for deterministic host integrations. Run it through ``uv run`` so the declared
``python-docx`` dependency is available without modifying the user's Python.
"""

from __future__ import annotations

import argparse
import copy
import json
import re
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ModuleNotFoundError as exc:  # pragma: no cover - runtime preflight
    raise SystemExit(
        "python-docx is required. Run this script with `uv run scripts/sow_build.py ...` "
        "or install the dependency from scripts/requirements-sow.txt."
    ) from exc


PLUGIN_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SPEC_PATH = PLUGIN_ROOT / "references" / "sow-format-default.md"
PAGE_TOKEN = "\u0000PAGE_FIELD\u0000"


def _parse_spec_from_markdown(md_text: str) -> dict[str, Any] | None:
    """Extract the first fenced JSON format-spec block from template Markdown."""
    match = re.search(r"```json\s*\n(.*?)\n```", md_text, re.S)
    if not match:
        return None
    parsed = json.loads(match.group(1))
    if not isinstance(parsed, dict):
        raise ValueError("SOW format spec must be a JSON object")
    return parsed


def _read_spec(path: Path) -> dict[str, Any]:
    parsed = _parse_spec_from_markdown(path.read_text(encoding="utf-8"))
    if parsed is None:
        raise ValueError(f"No fenced JSON format spec found in {path}")
    return parsed


def _deep_merge(base: dict[str, Any], overlay: dict[str, Any]) -> dict[str, Any]:
    result = copy.deepcopy(base)
    for key, value in overlay.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = _deep_merge(result[key], value)
        else:
            result[key] = copy.deepcopy(value)
    return result


def load_spec(
    spec_path: str | Path | None = None,
    default_path: str | Path = DEFAULT_SPEC_PATH,
) -> dict[str, Any]:
    """Load the packaged default, then overlay the user's derived format spec."""
    default = _read_spec(Path(default_path).expanduser().resolve())
    if not spec_path:
        return default
    path = Path(spec_path).expanduser().resolve()
    if not path.exists():
        return default
    return _deep_merge(default, _read_spec(path))


def _rgb(hexstr: str | None) -> RGBColor | None:
    if not hexstr:
        return None
    value = hexstr.lstrip("#")
    if len(value) != 6:
        raise ValueError(f"Expected six-digit RGB color, got {hexstr!r}")
    return RGBColor.from_string(value.upper())


def _alignment(value: Any) -> WD_ALIGN_PARAGRAPH | None:
    if value is None:
        return None
    if isinstance(value, int):
        return WD_ALIGN_PARAGRAPH(value)
    normalized = str(value).upper()
    for label, member in {
        "LEFT": WD_ALIGN_PARAGRAPH.LEFT,
        "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
        "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT,
        "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.items():
        if label in normalized:
            return member
    return None


def _set_cell_shading(cell: Any, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill.lstrip("#"))


def _set_cell_margins(cell: Any, margins: dict[str, float]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "left", "bottom", "right"):
        if edge not in margins:
            continue
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(int(float(margins[edge]) * 1440)))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table: Any, border_spec: dict[str, Any] | str | None) -> None:
    if not border_spec:
        return
    if isinstance(border_spec, str):
        border_spec = {"style": border_spec}
    style = str(border_spec.get("style", "single"))
    if style in {"none", "nil", "borderless"}:
        style = "nil"
    color = str(border_spec.get("color", "auto")).lstrip("#")
    size = str(border_spec.get("size", 4))
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), style)
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def _set_col_widths(table: Any, widths_in: list[float]) -> None:
    if len(widths_in) != len(table.columns):
        raise ValueError(
            f"Expected {len(table.columns)} column widths, got {len(widths_in)}"
        )
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gridcol, width in zip(grid.findall(qn("w:gridCol")), widths_in):
            gridcol.set(qn("w:w"), str(int(float(width) * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            cell.width = Inches(float(width))


def _repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def _suppress_paragraph_borders(paragraph: Any) -> None:
    """Override borders inherited from Word's built-in Title style."""
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        paragraph_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


class _SafeFormat(dict[str, str]):
    def __missing__(self, key: str) -> str:
        return "{" + key + "}"


class SOWDoc:
    def __init__(
        self,
        spec_path: str | Path | None = None,
        *,
        default_path: str | Path = DEFAULT_SPEC_PATH,
        client: str = "Client",
        vendor: str = "Vendor",
    ) -> None:
        self.spec = load_spec(spec_path, default_path)
        self.client = client
        self.vendor = vendor
        self.doc = Document()
        section = self.doc.sections[0]
        page = self.spec["page"]
        section.page_width = Inches(page["width_in"])
        section.page_height = Inches(page["height_in"])
        section.top_margin = Inches(page["margin_top_in"])
        section.bottom_margin = Inches(page["margin_bottom_in"])
        section.left_margin = Inches(page["margin_left_in"])
        section.right_margin = Inches(page["margin_right_in"])
        self._section_num = 0
        self._setup_footer(section)

    def _role(self, name: str) -> dict[str, Any]:
        roles = self.spec["roles"]
        return roles.get(name, roles["body"])

    def _apply_run_role(self, run: Any, role_name: str) -> None:
        role = self._role(role_name)
        font_name = role.get("name")
        if font_name:
            run.font.name = font_name
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
        if role.get("size_pt") is not None:
            run.font.size = Pt(role["size_pt"])
        if role.get("bold") is not None:
            run.font.bold = bool(role["bold"])
        if role.get("italic") is not None:
            run.font.italic = bool(role["italic"])
        color = _rgb(role.get("color"))
        if color is not None:
            run.font.color.rgb = color

    def _apply_paragraph_role(self, paragraph: Any, role_name: str) -> None:
        role = self._role(role_name)
        before = role.get("space_before_pt")
        after = role.get("space_after_pt")
        if before is not None:
            paragraph.paragraph_format.space_before = Pt(before)
        if after is not None:
            paragraph.paragraph_format.space_after = Pt(after)
        line_spacing = role.get("line_spacing")
        if line_spacing is not None:
            paragraph.paragraph_format.line_spacing = line_spacing
        alignment = _alignment(role.get("alignment"))
        if alignment is not None:
            paragraph.alignment = alignment

    def _footer_run(self, paragraph: Any, text: str) -> None:
        if not text:
            return
        run = paragraph.add_run(text)
        self._apply_run_role(run, "footer")

    def _setup_footer(self, section: Any) -> None:
        footer = section.footer
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.alignment = _alignment(self._role("footer").get("alignment")) or WD_ALIGN_PARAGRAPH.CENTER
        pattern = str(self.spec.get("footer_text", ""))
        values = _SafeFormat(client=self.client, vendor=self.vendor, page=PAGE_TOKEN)
        rendered = pattern.format_map(values)
        parts = rendered.split(PAGE_TOKEN)
        for index, part in enumerate(parts):
            self._footer_run(paragraph, part)
            if index < len(parts) - 1:
                _add_page_field(paragraph)
        if PAGE_TOKEN not in rendered and self.spec.get("footer_has_page_field", True):
            self._footer_run(paragraph, "  Page ")
            _add_page_field(paragraph)

    def title_block(self, title: str, subtitle: str | None = None, draft: bool = False) -> None:
        if draft or self.spec.get("draft_banner"):
            paragraph = self.doc.add_paragraph()
            self._apply_paragraph_role(paragraph, "draft_banner")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run("DRAFT FOR REVIEW - NOT EXECUTED")
            self._apply_run_role(run, "draft_banner")
        paragraph = self.doc.add_paragraph()
        paragraph.style = self.doc.styles["Title"]
        _suppress_paragraph_borders(paragraph)
        self._apply_paragraph_role(paragraph, "doc_title")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        self._apply_run_role(run, "doc_title")
        if subtitle:
            subtitle_paragraph = self.doc.add_paragraph()
            subtitle_paragraph.style = self.doc.styles["Subtitle"]
            self._apply_paragraph_role(subtitle_paragraph, "title")
            subtitle_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle_run = subtitle_paragraph.add_run(subtitle)
            self._apply_run_role(subtitle_run, "title")
        rule_color = self.spec.get("colors", {}).get("rule")
        if rule_color:
            self._add_horizontal_rule(rule_color)

    def _add_horizontal_rule(self, hex_color: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph_properties = paragraph._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(self.spec.get("rule_size", 12)))
        bottom.set(qn("w:color"), hex_color.lstrip("#"))
        border.append(bottom)
        paragraph_properties.append(border)

    def heading(self, _level: int, text: str) -> int:
        self._section_num += 1
        paragraph = self.doc.add_paragraph()
        paragraph.style = self.doc.styles["Heading 1"]
        self._apply_paragraph_role(paragraph, "section_heading")
        paragraph.paragraph_format.keep_with_next = True
        number_format = str(self.spec.get("heading_number_format", "{n}.  {title}"))
        run = paragraph.add_run(number_format.format(n=self._section_num, title=text))
        self._apply_run_role(run, "section_heading")
        return self._section_num

    def _add_markdown_bold_runs(self, paragraph: Any, text: str, role_name: str) -> None:
        for part in re.split(r"(\*\*.*?\*\*)", text):
            if not part:
                continue
            bold = part.startswith("**") and part.endswith("**")
            run = paragraph.add_run(part[2:-2] if bold else part)
            self._apply_run_role(run, role_name)
            if bold:
                run.font.bold = True

    def para(self, text: str) -> Any:
        paragraph = self.doc.add_paragraph()
        self._apply_paragraph_role(paragraph, "body")
        self._add_markdown_bold_runs(paragraph, text, "body")
        return paragraph

    def bullets(self, items: list[str]) -> None:
        indent = float(self.spec.get("bullet_indent_in", 0.25))
        for item in items:
            paragraph = self.doc.add_paragraph(style=None)
            self._apply_paragraph_role(paragraph, "body")
            paragraph.paragraph_format.left_indent = Inches(indent)
            bullet = paragraph.add_run("-  ")
            self._apply_run_role(bullet, "body")
            self._add_markdown_bold_runs(paragraph, item, "body")

    def table(
        self,
        header: list[str],
        rows: list[list[Any]],
        widths: list[float] | None = None,
    ) -> Any:
        if not header:
            raise ValueError("A table must have at least one header column")
        table = self.doc.add_table(rows=1, cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_spec = (self.spec.get("tables") or [{}])[0]
        header_fill = self.spec.get("colors", {}).get("table_header_fill") or table_spec.get("header_fill")
        cell_margins = table_spec.get("cell_margins_in", {})
        for index, value in enumerate(header):
            cell = table.rows[0].cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].text = ""
            self._apply_paragraph_role(cell.paragraphs[0], "table_header")
            run = cell.paragraphs[0].add_run(str(value))
            self._apply_run_role(run, "table_header")
            if header_fill:
                _set_cell_shading(cell, str(header_fill))
            _set_cell_margins(cell, cell_margins)
        _repeat_table_header(table.rows[0])
        for row in rows:
            if len(row) != len(header):
                raise ValueError(f"Table row has {len(row)} cells; expected {len(header)}")
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cells[index].paragraphs[0].text = ""
                self._apply_paragraph_role(cells[index].paragraphs[0], "table_body")
                run = cells[index].paragraphs[0].add_run(str(value))
                self._apply_run_role(run, "table_body")
                _set_cell_margins(cells[index], cell_margins)
        resolved_widths = widths or table_spec.get("col_widths_in")
        if resolved_widths and len(resolved_widths) == len(header):
            _set_col_widths(table, [float(value) for value in resolved_widths])
        _set_table_borders(table, table_spec.get("borders"))
        return table

    def signature(self, party_a: str, party_b: str) -> Any:
        layout = str(self.spec.get("signature_layout", "two_column_borderless"))
        if layout.startswith("one_column"):
            table = self.doc.add_table(rows=6, cols=1)
            labels = [
                party_a,
                "Signature: _______________________",
                "Name / Title / Date",
                party_b,
                "Signature: _______________________",
                "Name / Title / Date",
            ]
            for index, value in enumerate(labels):
                cell = table.rows[index].cells[0]
                cell.paragraphs[0].text = ""
                run = cell.paragraphs[0].add_run(value)
                self._apply_run_role(run, "signature_caption")
                if index in {0, 3}:
                    run.font.bold = True
        else:
            table = self.doc.add_table(rows=3, cols=2)
            signature_widths = self.spec.get("signature_col_widths_in", [3.25, 3.25])
            _set_col_widths(table, [float(value) for value in signature_widths])
            labels = [
                (party_a, party_b),
                ("Signature: _______________________", "Signature: _______________________"),
                ("Name / Title / Date", "Name / Title / Date"),
            ]
            for row_index, values in enumerate(labels):
                for column_index, value in enumerate(values):
                    cell = table.rows[row_index].cells[column_index]
                    cell.paragraphs[0].text = ""
                    run = cell.paragraphs[0].add_run(value)
                    self._apply_run_role(run, "signature_caption")
                    if row_index == 0:
                        run.font.bold = True
        _set_table_borders(table, "none" if "borderless" in layout else self.spec.get("signature_borders"))
        return table

    def exhibit(self, title: str) -> Any:
        self.doc.add_page_break()
        paragraph = self.doc.add_paragraph()
        self._apply_paragraph_role(paragraph, "section_heading")
        run = paragraph.add_run(title)
        self._apply_run_role(run, "section_heading")
        return paragraph

    def add_section(self, section: dict[str, Any]) -> None:
        if section.get("page_break_before"):
            self.doc.add_page_break()
        self.heading(int(section.get("level", 1)), str(section["heading"]))
        for block in section.get("blocks", []):
            kind = block.get("type")
            if kind == "paragraph":
                self.para(str(block.get("text", "")))
            elif kind == "bullets":
                self.bullets([str(item) for item in block.get("items", [])])
            elif kind == "table":
                self.table(
                    [str(item) for item in block.get("header", [])],
                    block.get("rows", []),
                    block.get("widths"),
                )
            elif kind == "page_break":
                self.doc.add_page_break()
            else:
                raise ValueError(f"Unsupported SOW block type: {kind!r}")

    def save(self, path: str | Path) -> Path:
        destination = Path(path).expanduser().resolve()
        destination.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(destination)
        return destination


def build_from_content(
    content: dict[str, Any],
    output_path: str | Path,
    *,
    spec_path: str | Path | None = None,
) -> Path:
    client = str(content.get("client", "Client"))
    vendor = str(content.get("vendor", "Vendor"))
    document = SOWDoc(spec_path, client=client, vendor=vendor)
    document.title_block(
        str(content.get("title", "Statement of Work")),
        str(content.get("subtitle", f"{client} x {vendor}")),
        bool(content.get("draft", True)),
    )
    for section in content.get("sections", []):
        document.add_section(section)
    signature = content.get("signature")
    if signature:
        document.signature(
            str(signature.get("party_a", client)),
            str(signature.get("party_b", vendor)),
        )
    return document.save(output_path)


def _smoke_content(client: str, vendor: str) -> dict[str, Any]:
    return {
        "title": "Statement of Work",
        "subtitle": f"{client} x {vendor}",
        "client": client,
        "vendor": vendor,
        "draft": True,
        "sections": [
            {
                "heading": "Relationship of the Parties",
                "blocks": [{"type": "paragraph", "text": "This is a **smoke test** render of the SOW builder."}],
            },
            {
                "heading": "Compensation",
                "blocks": [
                    {
                        "type": "table",
                        "header": ["Phase", "Not-to-Exceed"],
                        "rows": [["Phase 1", "$24,000"], ["Phase 2", "$50,000"]],
                        "widths": [3, 2.5],
                    }
                ],
            },
        ],
        "signature": {"party_a": client, "party_b": vendor},
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec_path", nargs="?", help="clients.sow-template.md path")
    parser.add_argument("output_path", nargs="?", help="destination .docx path")
    parser.add_argument("--content", help="structured SOW content JSON")
    parser.add_argument("--client", default="Sample Client Inc.")
    parser.add_argument("--vendor", default="BrightWay AI")
    parser.add_argument("--check", action="store_true", help="validate runtime and packaged default spec")
    args = parser.parse_args(argv)

    load_spec(args.spec_path)
    if args.check:
        print("SOW runtime ready: python-docx available and format spec valid.")
        return 0

    output = args.output_path or "/tmp/sow_smoke_test.docx"
    if args.content:
        content = json.loads(Path(args.content).read_text(encoding="utf-8"))
        if not isinstance(content, dict):
            raise ValueError("SOW content JSON must be an object")
    else:
        content = _smoke_content(args.client, args.vendor)
    destination = build_from_content(content, output, spec_path=args.spec_path)
    print(f"wrote {destination}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
